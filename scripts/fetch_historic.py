"""One-off script: download and parse historic season data for 2021-2025.

Writes tables/historic/{year}/combined.csv for each year.
Raw Excel downloads are saved to downloads/historic/{year}/ (gitignored).

Key differences from the current-season pipeline:
  - Historic Results XLS files already contain full team names — no abbreviation step.
  - Historic Matrix XLSX files use some abbreviation prefixes that differ from the
    current season (GOSL vs GOSLG, WYMO vs WYMOY, etc.) — handled via HISTORIC_ABBREVS.
  - 2023 files don't exist on the server; 2025 has no year-specific schedule.
"""

import csv
import pathlib
import requests
import xlrd
import openpyxl
from datetime import datetime

from parse import PROJECT_ROOT
from parse import expand as _expand_current
from parse.results import COLUMNS_NEEDED, RENAME
from parse.combine import RESULT_FIELDS
from parse.leagues import leagues as _leagues

YEARS = range(2021, 2026)
DOWNLOAD_DIR = PROJECT_ROOT / "downloads" / "historic"
HISTORIC_TABLES_DIR = PROJECT_ROOT / "tables" / "historic"

# Historic results files store full club names but sometimes use alias spellings
# that differ from the canonical name in TEAM_ABBREVIATIONS.
HISTORIC_NAME_ALIASES = {
    "WHEATHAM":  "WHEATHAMPSTEAD",   # alias used in older results XLS
    "ST MARGS":  "ST MARG",          # alias used in older results XLS
}


def _normalize(full_name):
    """Canonicalise a full team name from a historic results file."""
    full_name = full_name.strip()
    parts = full_name.rsplit(" ", 1)
    if len(parts) == 2 and parts[1].isdigit():
        club = HISTORIC_NAME_ALIASES.get(parts[0], parts[0])
        return club + " " + parts[1]
    return HISTORIC_NAME_ALIASES.get(full_name, full_name)


# Abbreviation prefixes used in historic Matrix files that differ from the current season.
# None means the team is a placeholder/unknown and should be skipped.
HISTORIC_ABBREVS = {
    "DATC":  "DATCHWORTH",   # 2021  (current season uses DATCH)
    "HBST":  "HBSTC",        # 2021  (club no longer in the league)
    "GOSL":  "GOSLING",      # 2022  (current uses GOSLG)
    "HODDN": "HODDESDON",    # 2022-2024  (current uses HODD)
    "LISTR": "LISTER",       # 2022  (current uses LIST)
    "WYMO":  "WYMONDLEY",    # 2021  (current uses WYMOY)
    "WESTN": "WESTON",       # 2024  (current uses WEST)
    "ST P":  "ST PAULS",     # 2021-2022  (current uses ST PS)
    "XXXX":  None,           # fixture-grid placeholders — skip
    "SPAR":  None,           # club never appeared in results — skip
}


def _expand(code):
    """Expand a Matrix abbreviation code to a full club name.

    Returns None for known placeholder codes that should be skipped.
    """
    stripped = code.strip()
    i = len(stripped)
    while i > 0 and stripped[i - 1].isdigit():
        i -= 1
    prefix = stripped[:i].rstrip()
    number = stripped[i:]

    if prefix in HISTORIC_ABBREVS:
        full = HISTORIC_ABBREVS[prefix]
        if full is None:
            return None
        return full + (" " + number if number else "")

    try:
        return _expand_current(stripped)
    except KeyError:
        print(f"  WARNING: unknown schedule code '{stripped}', skipping")
        return None


def _download(url, dest):
    if dest.exists():
        print(f"  already have {dest.name}")
        return True
    r = requests.get(url)
    if r.status_code == 404:
        print(f"  404: {url}")
        return False
    r.raise_for_status()
    dest.write_bytes(r.content)
    print(f"  downloaded {dest.name}")
    return True


def _parse_results(xls_path):
    wb = xlrd.open_workbook(str(xls_path))
    ws = wb.sheet_by_index(0)
    headers = ws.row_values(0)
    col_index = {name: headers.index(name) for name in COLUMNS_NEEDED}

    rows = []
    for row_idx in range(1, ws.nrows):
        raw = ws.row_values(row_idx)
        if raw[col_index["HomeAway"]] != "Home":
            continue
        record = {RENAME.get(col, col): raw[col_index[col]]
                  for col in COLUMNS_NEEDED if col != "HomeAway"}
        # Historic results store full names but may use alias spellings — normalise.
        record["home_team"]    = _normalize(record["home_team"])
        record["away_team"]    = _normalize(record["away_team"])
        record["fixture_date"] = xlrd.xldate_as_datetime(record["fixture_date"], wb.datemode).date()
        record["division"]     = record["division"].lstrip("0") or "0"
        record["home_sets"]    = int(record["home_sets"])
        record["away_sets"]    = int(record["away_sets"])
        record["home_games"]   = int(record["home_games"])
        record["away_games"]   = int(record["away_games"])
        rows.append(record)
    return rows


def _parse_schedule(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))

    divisions = []
    for col_idx in range(1, len(rows[1]), 2):
        div_name = rows[1][col_idx]
        if div_name is not None:
            divisions.append((str(div_name), col_idx))

    fixtures = []
    data_rows = rows[3:]
    for group_start in range(0, len(data_rows), 4):
        group = data_rows[group_start:group_start + 4]
        fixture_date = group[0][0]
        if fixture_date is None:
            continue
        if isinstance(fixture_date, datetime):
            fixture_date = fixture_date.date()
        for div_name, col_idx in divisions:
            for row in group:
                home_raw = row[col_idx]     if col_idx     < len(row) else None
                away_raw = row[col_idx + 1] if col_idx + 1 < len(row) else None
                if home_raw is None or away_raw is None:
                    continue
                home = _expand(str(home_raw))
                away = _expand(str(away_raw))
                if home is None or away is None:
                    continue
                fixtures.append({
                    "fixture_date": fixture_date,
                    "division":     div_name,
                    "home_team":    home,
                    "away_team":    away,
                })
    return fixtures


def _combine_and_write(schedule_rows, results_rows, out_path):
    results_index = {
        (str(r["fixture_date"]), r["division"], r["home_team"], r["away_team"]): r
        for r in results_rows
    }

    combined = []
    for row in schedule_rows:
        key = (str(row["fixture_date"]), row["division"], row["home_team"], row["away_team"])
        result = results_index.get(key, {})
        combined.append({
            **row,
            **{field: result.get(field, "") for field in RESULT_FIELDS},
        })

    out_path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(schedule_rows[0].keys()) + RESULT_FIELDS
    with open(out_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(combined)

    matched = sum(1 for r in combined if r["home_sets"] != "")
    print(f"  {len(combined)} fixtures, {matched} with results -> {out_path.relative_to(PROJECT_ROOT)}")


def _schedule_from_results(results_rows):
    """Synthesise a schedule from results when no Matrix file is available.

    Each result row already has date/division/teams so we can use it directly
    as the schedule fixture.  The combined CSV will have a result for every row.
    """
    seen = set()
    fixtures = []
    for r in results_rows:
        key = (str(r["fixture_date"]), r["division"], r["home_team"], r["away_team"])
        if key not in seen:
            seen.add(key)
            fixtures.append({
                "fixture_date": r["fixture_date"],
                "division":     r["division"],
                "home_team":    r["home_team"],
                "away_team":    r["away_team"],
            })
    return fixtures


if __name__ == "__main__":
    for year in YEARS:
        print(f"\n=== {year} ===")
        try:
            year_dl = DOWNLOAD_DIR / str(year)
            year_dl.mkdir(parents=True, exist_ok=True)

            results_path  = year_dl / f"{year}Results.xls"
            schedule_path = year_dl / f"{year}Matrix.xlsx"

            if not _download(f"https://www.datchworth.net/images/{year}Results.xls", results_path):
                print("  skipping year (results not available)")
                continue

            results_rows = _parse_results(results_path)

            schedule_available = _download(
                f"https://www.datchworth.net/images/{year}Matrix.xlsx", schedule_path
            )
            if schedule_available:
                schedule_rows = _parse_schedule(schedule_path)
            else:
                print("  no Matrix file — synthesising schedule from results")
                schedule_rows = _schedule_from_results(results_rows)

            print(f"  parsed {len(results_rows)} result rows, {len(schedule_rows)} schedule fixtures")

            out_path = HISTORIC_TABLES_DIR / str(year) / "combined.csv"
            _combine_and_write(schedule_rows, results_rows, out_path)
            _leagues(out_path, out_path.parent / "leagues.csv")

        except Exception as e:
            import traceback
            traceback.print_exc()

    print("\nDone.")
